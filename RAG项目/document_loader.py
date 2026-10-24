import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ALLOWED_EXTENSIONS = {".txt", ".docx", ".pdf"}
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


class DocumentParseError(ValueError):
    """文档无法解析或内容为空。"""


def extract_document_text(file_name: str, data: bytes) -> str:
    """从 TXT、DOCX 或 PDF 文件中提取纯文本。"""
    suffix = Path(file_name).suffix.lower()

    if suffix not in ALLOWED_EXTENSIONS:
        raise DocumentParseError("仅支持 TXT、DOCX 和 PDF 文件")
    if not data:
        raise DocumentParseError("文件内容为空")
    if len(data) > MAX_FILE_SIZE:
        raise DocumentParseError("文件不能超过 10 MB")

    try:
        if suffix == ".txt":
            text = _extract_txt(data)
        elif suffix == ".docx":
            text = _extract_docx(data)
        else:
            text = _extract_pdf(data)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(
            f"文件解析失败，请确认文件未损坏且格式正确：{exc}"
        ) from exc

    text = _normalize_text(text)
    if not text:
        if suffix == ".pdf":
            raise DocumentParseError(
                "PDF 中没有可提取的文字。扫描版 PDF 需要先进行 OCR。"
            )
        raise DocumentParseError("文档中没有可入库的文字内容")
    return text


def _extract_txt(data: bytes) -> str:
    """兼容 UTF-8 与常见中文编码的 TXT 文件。"""
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("TXT 文件编码无法识别，请转为 UTF-8 后重试")


def _extract_docx(data: bytes) -> str:
    """读取 Word 正文段落和表格，表格按行转换为制表符分隔文本。"""
    document = Document(io.BytesIO(data))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            row_text = "\t".join(cells).strip()
            if row_text:
                blocks.append(row_text)

    return "\n".join(block for block in blocks if block)


def _extract_pdf(data: bytes) -> str:
    """按页读取 PDF 文本，并保留页码方便检索结果定位。"""
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentParseError("PDF 已加密，无法读取内容")

    pages = []
    for page_number, page in enumerate(reader.pages, start=1):
        page_text = (page.extract_text() or "").strip()
        if page_text:
            pages.append(f"[第 {page_number} 页]\n{page_text}")
    return "\n\n".join(pages)


def _normalize_text(text: str) -> str:
    """统一换行并清理多余空白，保留段落结构。"""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

